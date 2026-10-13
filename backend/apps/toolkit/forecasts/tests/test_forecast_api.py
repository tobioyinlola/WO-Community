import io

import pytest

from apps.toolkit.ai.providers.fake import FakeLLMProvider
from apps.toolkit.ai.research.base import Page, SearchHit
from apps.toolkit.ai.research.fake import FakeResearchProvider
from apps.toolkit.forecasts import engine
from apps.toolkit.forecasts.models import Forecast
from apps.toolkit.testing import BASE
from apps.toolkit.ventures.models import OutputState

pytestmark = pytest.mark.django_db


def url(venture, tail=""):
    return f"{BASE}/ventures/{venture['id']}/forecast{tail}"


def started(client, venture):
    response = client.post(url(venture), {}, format="json")
    assert response.status_code == 201, response.content
    return response.json()


def edited(forecast, **drivers):
    a = {**forecast["assumptions"]}
    a["drivers"] = {**a["drivers"], **drivers}
    return a


def test_it_starts_from_the_profile(client, full_venture):
    forecast = started(client, full_venture)
    assert forecast["version"] == 1
    assert forecast["template"] == "transaction_fee"  # the first revenue model entered
    assert forecast["currency"] == "NGN"  # from the price point
    state = OutputState.objects.get(venture_id=full_venture["id"], output_type="forecast")
    assert state.status == "ready"
    again = client.post(url(full_venture), {}, format="json")
    assert again.status_code == 409


def test_saving_checks_every_value_and_the_version(client, full_venture):
    forecast = started(client, full_venture)
    ok = client.put(
        url(full_venture),
        {"assumptions": edited(forecast, avg_value_minor=500000), "base_version": 1},
        format="json",
    )
    assert ok.status_code == 200 and ok.json()["version"] == 2
    stale = client.put(
        url(full_venture),
        {"assumptions": edited(forecast, avg_value_minor=1), "base_version": 1},
        format="json",
    )
    assert stale.status_code == 409
    bad = client.put(
        url(full_venture),
        {"assumptions": edited(forecast, fee_pct=500), "base_version": 2},
        format="json",
    )
    assert bad.status_code == 400
    assert "drivers.fee_pct" in str(bad.json())
    assert client.get(url(full_venture)).json()["version"] == 2


def test_the_result_is_24_months_of_integer_minor_units(client, full_venture):
    forecast = started(client, full_venture)
    client.put(
        url(full_venture),
        {
            "assumptions": edited(
                forecast, avg_value_minor=500000, transactions_month1=400, fee_pct=2
            ),
            "base_version": 1,
        },
        format="json",
    )
    body = client.get(url(full_venture, "/result")).json()
    assert len(body["months"]) == 24 and body["scenario"] == "base"
    first = body["months"][0]
    assert first["revenue"] == 400 * 10000  # 400 transactions of 5,000.00 at 2%
    assert all(isinstance(m["cash_end"], int) for m in body["months"])
    low = client.get(url(full_venture, "/result"), {"scenario": "conservative"}).json()
    assert low["summary"]["revenue_year2"] < body["summary"]["revenue_year2"]
    assert client.get(url(full_venture, "/result"), {"scenario": "wild"}).status_code == 400


def test_scenarios_and_sensitivity(client, full_venture):
    forecast = started(client, full_venture)
    client.put(
        url(full_venture),
        {"assumptions": edited(forecast, avg_value_minor=500000), "base_version": 1},
        format="json",
    )
    scenarios = client.get(url(full_venture, "/scenarios")).json()
    assert (
        scenarios["conservative"]["revenue_year2"]
        < scenarios["base"]["revenue_year2"]
        < scenarios["optimistic"]["revenue_year2"]
    )
    rows = client.get(url(full_venture, "/sensitivity")).json()
    assert rows and rows[0]["impact"] >= rows[-1]["impact"]


def test_versions_can_be_restored(client, full_venture):
    forecast = started(client, full_venture)
    client.put(
        url(full_venture),
        {"assumptions": edited(forecast, fee_pct=9), "base_version": 1},
        format="json",
    )
    versions = client.get(url(full_venture, "/versions")).json()
    assert [v["version"] for v in versions] == [2, 1]
    restored = client.post(url(full_venture, "/versions/1/restore"))
    assert restored.status_code == 200
    body = restored.json()
    assert body["version"] == 3 and body["assumptions"]["drivers"]["fee_pct"] == 1.5
    assert client.post(url(full_venture, "/versions/9/restore")).status_code == 404


def test_a_changed_price_marks_the_forecast_out_of_date(client, full_venture):
    started(client, full_venture)
    client.patch(
        f"{BASE}/ventures/{full_venture['id']}/profile",
        {
            "fields": {
                "price_points": [
                    {"label": "Pro", "amount_minor": 900000, "currency": "NGN", "period": "month"}
                ]
            }
        },
        format="json",
    )
    package = client.get(f"{BASE}/ventures/{full_venture['id']}/package").json()
    state = next(o for o in package["outputs"] if o["type"] == "forecast")
    assert state["status"] == "needs_refresh" and "price_points" in state["stale_fields"]


def research():
    page = Page(
        url="https://data.example.org/saas",
        title="Benchmarks",
        text="Median monthly churn for small business software is 4 percent.",
        publisher="Example Data",
    )
    for key in ("customer acquisition", "growth"):
        FakeResearchProvider.hits[key] = [SearchHit(page.url, page.title)]
    FakeResearchProvider.pages[page.url] = page
    FakeLLMProvider.canned["research.extract_facts"] = lambda v, m: {
        "facts": (
            [
                {
                    "statement": "Median monthly churn for small business software is 4 percent",
                    "quote": "Median monthly churn for small business software is 4 percent",
                    "value": 4,
                    "unit": "percent",
                    "region": "africa",
                    "confidence": 0.8,
                }
            ]
            if v["topic"] == "benchmarks"
            else []
        )
    }


def test_suggestions_are_sourced_proposals_and_save_nothing(client, full_venture):
    forecast = started(client, full_venture)
    research()
    FakeLLMProvider.canned["forecast.assumptions"] = {
        "proposals": [
            {"name": "transactions_growth_pct", "value": 6, "rationale": "Cautious.", "fact": 1},
            {"name": "fee_pct", "value": 900, "rationale": "Out of range.", "fact": 1},
            {"name": "avg_value_minor", "value": 1, "rationale": "A price.", "fact": 1},
            {
                "name": "transactions_growth_pct",
                "value": 5,
                "rationale": "No such fact.",
                "fact": 7,
            },
        ]
    }
    response = client.post(url(full_venture, "/suggest"))
    assert response.status_code == 200, response.content
    proposals = response.json()["proposals"]
    assert [p["name"] for p in proposals] == ["transactions_growth_pct"]
    assert proposals[0]["url"] == "https://data.example.org/saas"
    assert proposals[0]["current"] == forecast["assumptions"]["drivers"]["transactions_growth_pct"]
    assert Forecast.objects.get().version == 1  # nothing was written


def test_the_spreadsheet_keeps_its_formulas(client, full_venture, run_jobs):
    from openpyxl import load_workbook

    from apps.integrations.storage.base import PRIVATE
    from apps.integrations.storage.fake import FakeStorage
    from apps.toolkit.jobs.models import ExportFile

    forecast = started(client, full_venture)
    client.put(
        url(full_venture),
        {"assumptions": edited(forecast, avg_value_minor=500000), "base_version": 1},
        format="json",
    )
    job = client.post(url(full_venture, "/export"), {"format": "xlsx"}, format="json")
    assert job.status_code == 202
    run_jobs()
    done = client.get(f"{BASE}/jobs/{job.json()['id']}").json()
    assert done["status"] == "succeeded", done
    file = ExportFile.objects.get(pk=done["result"]["id"])
    book = load_workbook(io.BytesIO(FakeStorage.objects[(PRIVATE, file.storage_key)][0]))
    assert book.sheetnames[:3] == ["Assumptions", "Summary", "Base"]
    base = book["Base"]
    assert base["B4"].value == "=B2-B3"  # gross profit is a formula
    assert base["B13"].value == "=Assumptions!$B$6+B11+B12"
    assert base["C13"].value == "=B13+C11+C12"
    assert isinstance(base["B2"].value, int | float)
    pdf = client.post(url(full_venture, "/export"), {"format": "pdf"}, format="json")
    run_jobs()
    again = client.get(f"{BASE}/jobs/{pdf.json()['id']}").json()
    assert again["status"] == "succeeded"
    stored = ExportFile.objects.get(pk=again["result"]["id"])
    assert FakeStorage.objects[(PRIVATE, stored.storage_key)][0].startswith(b"%PDF")


def test_only_the_owner_reaches_a_forecast(client, other_client, full_venture):
    started(client, full_venture)
    for tail in ("", "/result", "/scenarios", "/sensitivity", "/versions"):
        assert other_client.get(url(full_venture, tail)).status_code == 404
    assert other_client.post(url(full_venture, "/suggest")).status_code == 404
    assert other_client.post(url(full_venture, "/export"), {}, format="json").status_code == 404


def test_anonymous_and_unapproved_members_are_turned_away(api_client, pending_client, full_venture):
    assert api_client.get(url(full_venture)).status_code == 401
    assert pending_client.get(url(full_venture)).status_code == 403
    assert pending_client.get(f"{BASE}/forecast-templates").status_code == 403


def test_the_flag_switches_the_tool_off(client, full_venture, settings):
    settings.TOOLKIT_FEATURES = {**settings.TOOLKIT_FEATURES, "forecast": False}
    assert client.get(url(full_venture)).status_code == 404
    assert client.get(f"{BASE}/forecast-templates").status_code == 404


def test_templates_list_every_template_with_its_drivers(client):
    body = client.get(f"{BASE}/forecast-templates").json()
    assert [t["template"] for t in body] == list(engine.TEMPLATES)
    assert all(t["drivers"] and t["label"] for t in body)
