import io

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from PIL import Image

from apps.integrations.storage.base import PRIVATE
from apps.integrations.storage.fake import FakeStorage
from apps.toolkit.ai.providers.fake import FakeLLMProvider
from apps.toolkit.decks.models import Deck, DeckAsset
from apps.toolkit.jobs.models import ExportFile
from apps.toolkit.testing import BASE

pytestmark = pytest.mark.django_db


def url(venture, tail=""):
    return f"{BASE}/ventures/{venture['id']}/decks{tail}"


def make(client, venture, template="investor_pitch"):
    response = client.post(url(venture), {"template": template}, format="json")
    assert response.status_code == 201, response.content
    return response.json()


def patch(client, venture, deck, ops, **extra):
    return client.patch(
        url(venture, f"/{deck['id']}"),
        {"base_version": deck["version"], "ops": ops, **extra},
        format="json",
    )


def png(size=(40, 30)):
    out = io.BytesIO()
    Image.new("RGB", size, (200, 30, 30)).save(out, format="PNG")
    return out.getvalue()


def test_templates_are_listed(client):
    body = client.get(f"{BASE}/deck-templates").json()
    assert [t["slug"] for t in body] == ["investor_pitch", "demo_day", "accelerator", "teaser"]
    assert body[0]["slide_count"] == 13


def test_a_deck_is_filled_from_the_profile_in_the_founders_own_words(client, full_venture):
    deck = make(client, full_venture)
    slides = deck["content"]["slides"]
    assert len(slides) == 13 and deck["version"] == 1
    problem = next(s for s in slides if s["title"] == "The problem")
    assert problem["elements"][0]["text"].startswith("Market traders lose sales")
    model = next(s for s in slides if s["title"] == "How we make money")
    assert model["elements"][0]["items"] == ["Transaction fees", "Subscriptions"]
    traction = next(s for s in slides if s["title"] == "Traction")
    assert traction["elements"] == []  # nothing entered, so nothing invented
    assert len({s["id"] for s in slides}) == 13


def test_an_unknown_template_is_refused(client, full_venture):
    response = client.post(url(full_venture), {"template": "nope"}, format="json")
    assert response.status_code == 400


def test_patch_operations_change_slides_and_are_validated(client, full_venture):
    deck = make(client, full_venture, "teaser")
    first = deck["content"]["slides"][1]
    ok = patch(
        client,
        full_venture,
        deck,
        [
            {"op": "replace", "path": "/slides/1/title", "value": "<b>The pain</b>"},
            {
                "op": "add",
                "path": "/slides/-",
                "value": {
                    "layout": "quote",
                    "title": "Voice",
                    "elements": [{"type": "text", "text": "<script>x()</script>Great."}],
                },
            },
            {"op": "move", "from": "/slides/5", "path": "/slides/0"},
        ],
    )
    assert ok.status_code == 200, ok.content
    body = ok.json()
    assert body["version"] == 2 and len(body["content"]["slides"]) == 6
    assert body["content"]["slides"][0]["layout"] == "quote"
    assert "<" not in str(body["content"])
    assert first["id"] in [s["id"] for s in body["content"]["slides"]]


@pytest.mark.parametrize(
    ("op", "field"),
    [
        ({"op": "add", "path": "/slides/-", "value": {"layout": "html", "elements": []}}, "layout"),
        ({"op": "replace", "path": "/profile", "value": 1}, "ops"),
        ({"op": "remove", "path": "/slides/99"}, "ops"),
        ({"op": "copy", "path": "/slides/0", "from": "/slides/1"}, "ops"),
        (
            {
                "op": "add",
                "path": "/slides/0/elements/-",
                "value": {"type": "image", "asset_id": "f" * 32, "alt": "x"},
            },
            "asset_id",
        ),
        (
            {
                "op": "add",
                "path": "/slides/0/elements/-",
                "value": {"type": "linked", "kind": "profile_field", "ref": "password"},
            },
            "ref",
        ),
        (
            {"op": "replace", "path": "/theme", "value": {"palette": "light", "accent": "red"}},
            "accent",
        ),
    ],
)
def test_bad_patches_are_refused_and_change_nothing(client, full_venture, op, field):
    deck = make(client, full_venture, "teaser")
    response = patch(client, full_venture, deck, [op])
    assert response.status_code == 400
    assert field in str(response.json())
    assert client.get(url(full_venture, f"/{deck['id']}")).json()["version"] == 1


def test_a_stale_version_is_a_conflict(client, full_venture):
    deck = make(client, full_venture, "teaser")
    op = [{"op": "replace", "path": "/slides/0/title", "value": "A"}]
    assert patch(client, full_venture, deck, op).status_code == 200
    assert patch(client, full_venture, deck, op).status_code == 409


def test_slide_and_element_limits(client, full_venture, settings):
    settings.DECK_MAX_SLIDES = 6
    deck = make(client, full_venture, "teaser")
    add = {"op": "add", "path": "/slides/-", "value": {"layout": "statement", "elements": []}}
    assert patch(client, full_venture, deck, [add]).status_code == 200
    again = client.get(url(full_venture, f"/{deck['id']}")).json()
    assert patch(client, full_venture, again, [add]).status_code == 400


def test_figures_are_linked_and_follow_the_forecast(client, full_venture):
    deck = make(client, full_venture)
    shown = client.get(url(full_venture, f"/{deck['id']}/shown")).json()
    money = next(
        e
        for s in shown["slides"]
        for e in s["elements"]
        if e["type"] == "linked" and e["kind"] == "forecast_metric" and e["ref"] == "revenue_year1"
    )
    assert money["resolved"]["missing"] is True

    created = client.post(
        f"{BASE}/ventures/{full_venture['id']}/forecast", {}, format="json"
    ).json()
    drivers = {**created["assumptions"]["drivers"], "avg_value_minor": 500000}
    client.put(
        f"{BASE}/ventures/{full_venture['id']}/forecast",
        {"assumptions": {**created["assumptions"], "drivers": drivers}, "base_version": 1},
        format="json",
    )
    shown = client.get(url(full_venture, f"/{deck['id']}/shown")).json()
    money = next(
        e
        for s in shown["slides"]
        for e in s["elements"]
        if e["type"] == "linked" and e["kind"] == "forecast_metric" and e["ref"] == "revenue_year1"
    )
    assert money["resolved"]["display"].startswith("NGN ")
    assert money["resolved"]["value_minor"] > 0
    chart = next(
        e["resolved"]
        for s in shown["slides"]
        for e in s["elements"]
        if e.get("kind") == "forecast_series"
    )
    assert len(chart["values"]) == 8 and chart["currency"] == "NGN"


def test_the_check_flags_gaps_and_unsupported_figures(client, full_venture):
    deck = make(client, full_venture, "teaser")
    patch(
        client,
        full_venture,
        deck,
        [
            {
                "op": "add",
                "path": "/slides/1/elements/-",
                "value": {"type": "text", "text": "We have 48213 paying customers."},
            }
        ],
    )
    findings = client.get(url(full_venture, f"/{deck['id']}/check")).json()
    codes = {f["code"] for f in findings}
    assert {"unsupported_figure", "no_notes"} <= codes
    figure = next(f for f in findings if f["code"] == "unsupported_figure")
    assert "48213" in figure["message"] and figure["slide"] == 2
    assert findings == sorted(
        findings, key=lambda f: ({"error": 0, "warning": 1, "info": 2}[f["severity"]], f["slide"])
    )


def test_assist_proposes_and_never_adds_a_figure(client, full_venture):
    deck = make(client, full_venture, "teaser")
    problem = deck["content"]["slides"][1]
    text_el = problem["elements"][0]
    FakeLLMProvider.canned["deck.slide"] = {
        "elements": [{"id": text_el["id"], "type": "text", "text": "Traders lose sales weekly."}]
    }
    ok = client.post(
        url(full_venture, f"/{deck['id']}/assist"),
        {"slide_id": problem["id"], "action": "shorten"},
        format="json",
    )
    assert ok.status_code == 200 and ok.json()["elements"][0]["text"].startswith("Traders lose")
    FakeLLMProvider.canned["deck.slide"] = {
        "elements": [
            {"id": text_el["id"], "type": "text", "text": "Traders lose 73000 sales weekly."}
        ]
    }
    bad = client.post(
        url(full_venture, f"/{deck['id']}/assist"),
        {"slide_id": problem["id"], "action": "punchier"},
        format="json",
    ).json()
    assert bad["elements"] == [] and bad["removed"] == ["73000"]
    saved = client.get(url(full_venture, f"/{deck['id']}")).json()
    assert saved["version"] == 1  # nothing was changed


def test_images_are_checked_stored_privately_and_used_by_id(client, full_venture):
    deck = make(client, full_venture, "teaser")
    upload = client.post(
        url(full_venture, f"/{deck['id']}/assets"),
        {"file": SimpleUploadedFile("logo.png", png(), content_type="image/png")},
        format="multipart",
    )
    assert upload.status_code == 201, upload.content
    asset = upload.json()
    stored = DeckAsset.objects.get(pk=asset["id"])
    assert FakeStorage.objects[(PRIVATE, stored.storage_key)][0][:4] == b"RIFF"  # re-encoded
    use = patch(
        client,
        full_venture,
        deck,
        [
            {
                "op": "add",
                "path": "/slides/0/elements/-",
                "value": {"type": "image", "asset_id": asset["id"], "alt": "Our logo"},
            }
        ],
    )
    assert use.status_code == 200, use.content
    in_use = client.delete(url(full_venture, f"/{deck['id']}/assets/{asset['id']}"))
    assert in_use.status_code == 409
    shown = client.get(url(full_venture, f"/{deck['id']}/shown")).json()
    image = next(e for e in shown["slides"][0]["elements"] if e["type"] == "image")
    assert image["url"]
    refused = client.post(
        url(full_venture, f"/{deck['id']}/assets"),
        {"file": SimpleUploadedFile("x.png", b"<svg onload=alert(1)>", content_type="image/png")},
        format="multipart",
    )
    assert refused.status_code == 400


def test_a_deck_exports_to_powerpoint_and_pdf(client, full_venture, run_jobs):
    from pptx import Presentation

    deck = make(client, full_venture, "demo_day")
    patch(
        client,
        full_venture,
        deck,
        [{"op": "replace", "path": "/slides/0/notes", "value": "Say hello"}],
    )
    started = client.post(
        url(full_venture, f"/{deck['id']}/export"), {"format": "pptx"}, format="json"
    )
    assert started.status_code == 202
    run_jobs()
    job = client.get(f"{BASE}/jobs/{started.json()['id']}").json()
    assert job["status"] == "succeeded", job
    file = ExportFile.objects.get(pk=job["result"]["id"])
    presentation = Presentation(io.BytesIO(FakeStorage.objects[(PRIVATE, file.storage_key)][0]))
    assert len(presentation.slides) == 8
    assert presentation.slides[0].notes_slide.notes_text_frame.text == "Say hello"

    again = client.post(
        url(full_venture, f"/{deck['id']}/export"), {"format": "pdf"}, format="json"
    )
    run_jobs()
    done = client.get(f"{BASE}/jobs/{again.json()['id']}").json()
    assert done["status"] == "succeeded"
    pdf = ExportFile.objects.get(pk=done["result"]["id"])
    assert FakeStorage.objects[(PRIVATE, pdf.storage_key)][0].startswith(b"%PDF")


def test_versions_can_be_restored(client, full_venture):
    deck = make(client, full_venture, "teaser")
    original = deck["content"]["slides"][0]["title"]
    changed = patch(
        client,
        full_venture,
        deck,
        [{"op": "replace", "path": "/slides/0/title", "value": "Changed"}],
        label="Before the meeting",
    ).json()
    versions = client.get(url(full_venture, f"/{deck['id']}/versions")).json()
    assert any(v["label"] == "Before the meeting" for v in versions)
    restored = client.post(url(full_venture, f"/{deck['id']}/versions/1/restore"))
    assert restored.status_code == 200
    assert restored.json()["content"]["slides"][0]["title"] == original
    assert restored.json()["version"] == changed["version"] + 1


def test_deleting_a_deck_removes_its_images(client, full_venture):
    deck = make(client, full_venture, "teaser")
    client.post(
        url(full_venture, f"/{deck['id']}/assets"),
        {"file": SimpleUploadedFile("logo.png", png(), content_type="image/png")},
        format="multipart",
    )
    assert DeckAsset.objects.count() == 1
    assert client.delete(url(full_venture, f"/{deck['id']}")).status_code == 204
    assert Deck.objects.count() == 0 and DeckAsset.objects.count() == 0
    assert not [k for k in FakeStorage.objects if k[0] == PRIVATE]


def test_only_the_owner_reaches_a_deck(client, other_client, full_venture):
    deck = make(client, full_venture, "teaser")
    base = url(full_venture, f"/{deck['id']}")
    for tail in ("", "/shown", "/check", "/versions", "/assets"):
        assert other_client.get(base + tail).status_code == 404
    assert other_client.get(url(full_venture)).status_code == 404
    assert (
        other_client.post(url(full_venture), {"template": "teaser"}, format="json").status_code
        == 404
    )
    assert other_client.delete(base).status_code == 404


def test_anonymous_and_unapproved_members_are_turned_away(api_client, pending_client, full_venture):
    assert api_client.get(url(full_venture)).status_code == 401
    assert pending_client.get(url(full_venture)).status_code == 403


def test_the_flag_switches_decks_off(client, full_venture, settings):
    settings.TOOLKIT_FEATURES = {**settings.TOOLKIT_FEATURES, "decks": False}
    assert client.get(url(full_venture)).status_code == 404
    assert client.get(f"{BASE}/deck-templates").status_code == 404
