"""A deck as a PowerPoint file or a PDF, from the resolved slides.

Both renderers use the same layout plan so they look alike: a title band, then the slide's
elements stacked in one or two columns. Colours and fonts come from the deck's named palette;
nothing the member typed is interpreted as markup.
"""

import io
from datetime import UTC, datetime
from typing import Any

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from reportlab.lib.colors import HexColor
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

from apps.toolkit.decks import schema

PPTX_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
PDF_TYPE = "application/pdf"
W, H = 13.333, 7.5  # inches
MARGIN = 0.7
TITLE_H = 1.1


def _palette(theme: dict[str, str]) -> dict[str, str]:
    colours = dict(schema.PALETTES.get(theme.get("palette", "light"), schema.PALETTES["light"]))
    if theme.get("accent"):
        colours["accent"] = theme["accent"].lstrip("#")
    return colours


def _plain(el: dict[str, Any]) -> list[tuple[str, str]]:
    """The element as (style, text) lines for a text frame: heading, body, bullet, metric."""
    kind = el["type"]
    if kind == "heading":
        return [("heading", el.get("text", ""))]
    if kind == "text":
        return [("body", el.get("text", ""))]
    if kind == "bullets":
        return [("bullet", i) for i in el.get("items", [])]
    if kind == "metric":
        lines = [("metric", f"{el.get('value', '')}"), ("small", el.get("label", ""))]
        if el.get("note"):
            lines.append(("small", el["note"]))
        return lines
    if kind == "table":
        lines = [("bold", " | ".join(el.get("columns", [])))]
        lines += [("body", " | ".join(r)) for r in el.get("rows", [])]
        return lines
    if kind == "linked":
        r = el.get("resolved", {})
        if el["kind"] == "profile_field":
            return [("body", r.get("text", ""))]
        if el["kind"] == "forecast_metric":
            if r.get("missing"):
                return [("small", r.get("reason", ""))]
            return [("metric", r["display"]), ("small", el.get("label", ""))]
        if el["kind"] == "forecast_series":
            return []  # drawn as a chart
        if el["kind"] == "team":
            return [
                ("bullet", f"{m['name']}, {m['role']}" if m["role"] else m["name"])
                for m in r.get("members", [])
            ]
        return [
            ("bullet", f"{c['name']}: {c['note']}" if c["note"] else c["name"])
            for c in r.get("competitors", [])
        ]
    return []


HEIGHTS = {"heading": 0.6, "body": 0.42, "bullet": 0.4, "metric": 0.8, "small": 0.3, "bold": 0.4}


def _estimate(lines: list[tuple[str, str]], width: float) -> float:
    per_line = max(20, int(width * 9))
    total = 0.0
    for style, text in lines:
        rows = max(1, -(-len(text) // per_line))
        total += HEIGHTS.get(style, 0.4) * (rows if style in ("body", "bullet") else 1)
    return total + 0.15


def plan(slide: dict[str, Any]) -> list[dict[str, Any]]:
    """Boxes (inches) for each element of a slide, in reading order."""
    body_top = MARGIN + TITLE_H if slide.get("title") else MARGIN + 1.0
    if slide["layout"] in ("title", "closing"):
        body_top = 3.0
    columns = 2 if slide["layout"] in ("two_column", "metrics", "team") else 1
    col_w = (W - 2 * MARGIN - (0.5 if columns == 2 else 0)) / columns
    ys = [body_top] * columns
    boxes: list[dict[str, Any]] = []
    for i, el in enumerate(slide["elements"]):
        col = i % columns if columns == 2 else 0
        chart = el["type"] == "linked" and el.get("kind") == "forecast_series"
        image = el["type"] == "image"
        lines = _plain(el)
        height = 3.0 if chart or image else _estimate(lines, col_w)
        if ys[col] + height > H - 0.4:
            ys[col] = min(ys[col], H - 0.4 - height)  # keep it on the page, never off it
        boxes.append(
            {
                "el": el,
                "x": MARGIN + col * (col_w + 0.5),
                "y": ys[col],
                "w": col_w,
                "h": height,
                "lines": lines,
                "chart": chart,
                "image": image,
            }
        )
        ys[col] += height + 0.15
    return boxes


def _rgb(hex6: str) -> RGBColor:
    return RGBColor.from_string(hex6.upper())


def render_pptx(title: str, deck: dict[str, Any], images: dict[str, bytes]) -> bytes:
    colours = _palette(deck.get("theme", {}))
    font = "Georgia" if deck.get("theme", {}).get("font") == "serif" else "Calibri"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    prs.core_properties.title = title[:120]
    prs.core_properties.created = datetime.now(UTC)
    blank = prs.slide_layouts[6]
    for slide in deck["slides"]:
        s = prs.slides.add_slide(blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = _rgb(colours["background"])
        if slide.get("title") and slide["layout"] not in ("title",):
            box = s.shapes.add_textbox(
                Inches(MARGIN), Inches(MARGIN), Inches(W - 2 * MARGIN), Inches(0.9)
            )
            run = box.text_frame.paragraphs[0].add_run()
            run.text = slide["title"]
            run.font.size, run.font.bold, run.font.name = Pt(34), True, font
            run.font.color.rgb = _rgb(colours["accent"])
        for b in plan(slide):
            if b["chart"]:
                r = b["el"].get("resolved", {})
                if r.get("missing"):
                    continue
                data = CategoryChartData()
                data.categories = r["labels"]
                data.add_series(b["el"].get("label") or "Series", [v / 100 for v in r["values"]])
                s.shapes.add_chart(
                    XL_CHART_TYPE.COLUMN_CLUSTERED,
                    Inches(b["x"]),
                    Inches(b["y"]),
                    Inches(b["w"]),
                    Inches(b["h"]),
                    data,
                )
                continue
            if b["image"]:
                blob = images.get(b["el"]["asset_id"])
                if blob:
                    s.shapes.add_picture(
                        io.BytesIO(blob), Inches(b["x"]), Inches(b["y"]), height=Inches(b["h"])
                    )
                continue
            if not b["lines"]:
                continue
            box = s.shapes.add_textbox(
                Inches(b["x"]), Inches(b["y"]), Inches(b["w"]), Inches(b["h"])
            )
            frame = box.text_frame
            frame.word_wrap = True
            for n, (style, text) in enumerate(b["lines"]):
                para = frame.paragraphs[0] if n == 0 else frame.add_paragraph()
                if slide["layout"] in ("title", "closing"):
                    para.alignment = PP_ALIGN.CENTER
                run = para.add_run()
                run.text = ("• " if style == "bullet" else "") + text
                run.font.name = font
                size = {"heading": 28, "metric": 36, "small": 14, "bold": 18}.get(style, 20)
                if slide["layout"] == "title" and n == 0:
                    size = 44
                run.font.size = Pt(size)
                run.font.bold = style in ("heading", "metric", "bold")
                run.font.color.rgb = _rgb(
                    colours["accent"]
                    if style == "metric"
                    else colours["muted"]
                    if style == "small"
                    else colours["text"]
                )
        if slide.get("notes"):
            s.notes_slide.notes_text_frame.text = slide["notes"]
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def render_pdf(title: str, deck: dict[str, Any], images: dict[str, bytes]) -> bytes:
    colours = _palette(deck.get("theme", {}))
    serif = deck.get("theme", {}).get("font") == "serif"
    regular, bold = ("Times-Roman", "Times-Bold") if serif else ("Helvetica", "Helvetica-Bold")
    out = io.BytesIO()
    pt = 72.0
    c = canvas.Canvas(out, pagesize=(W * pt, H * pt))
    c.setTitle(title[:120])
    c.setAuthor("")

    def y_of(inches_from_top: float) -> float:
        return (H - inches_from_top) * pt

    for slide in deck["slides"]:
        c.setFillColor(HexColor("#" + colours["background"]))
        c.rect(0, 0, W * pt, H * pt, stroke=0, fill=1)
        if slide.get("title") and slide["layout"] != "title":
            c.setFillColor(HexColor("#" + colours["accent"]))
            c.setFont(bold, 30)
            c.drawString(MARGIN * pt, y_of(MARGIN + 0.55), slide["title"])
        for b in plan(slide):
            if b["chart"]:
                r = b["el"].get("resolved", {})
                if r.get("missing"):
                    continue
                _bars(c, b, r, colours, regular, pt)
                continue
            if b["image"]:
                blob = images.get(b["el"]["asset_id"])
                if blob:
                    c.drawImage(
                        ImageReader(io.BytesIO(blob)),
                        b["x"] * pt,
                        y_of(b["y"] + b["h"]),
                        height=b["h"] * pt,
                        preserveAspectRatio=True,
                        mask="auto",
                    )
                continue
            y = b["y"]
            centred = slide["layout"] in ("title", "closing")
            for n, (style, text) in enumerate(b["lines"]):
                size = {"heading": 24, "metric": 32, "small": 12, "bold": 16}.get(style, 17)
                if slide["layout"] == "title" and n == 0:
                    size = 40
                font = bold if style in ("heading", "metric", "bold") or size >= 40 else regular
                color = (
                    colours["accent"]
                    if style == "metric"
                    else colours["muted"]
                    if style == "small"
                    else colours["text"]
                )
                c.setFillColor(HexColor("#" + color))
                c.setFont(font, size)
                shown = ("• " if style == "bullet" else "") + text
                lines = _wrap(shown, font, size, b["w"] * pt)
                for line in lines:
                    y += size * 1.35 / pt
                    if centred:
                        c.drawCentredString(W * pt / 2, y_of(y), line)
                    else:
                        c.drawString(b["x"] * pt, y_of(y), line)
        c.showPage()
    c.save()
    return out.getvalue()


def _wrap(text: str, font: str, size: float, width: float) -> list[str]:
    from reportlab.pdfbase.pdfmetrics import stringWidth

    words, lines, line = text.split(), [], ""
    for word in words:
        trial = f"{line} {word}".strip()
        if stringWidth(trial, font, size) <= width or not line:
            line = trial
        else:
            lines.append(line)
            line = word
    if line:
        lines.append(line)
    return lines or [""]


def _bars(
    c: Any, b: dict[str, Any], r: dict[str, Any], colours: dict[str, str], font: str, pt: float
) -> None:
    values = r["values"]
    top = max(max(values, default=0), 1)
    low = min(min(values, default=0), 0)
    span = top - low or 1
    x0, y0, w, h = b["x"] * pt, (H - b["y"] - b["h"]) * pt, b["w"] * pt, (b["h"] - 0.4) * pt
    gap = w / max(len(values), 1)
    c.setFont(font, 10)
    for i, v in enumerate(values):
        bar = (v - low) / span * h
        c.setFillColor(HexColor("#" + colours["accent"]))
        c.rect(x0 + i * gap + gap * 0.15, y0 + 0.4 * pt, gap * 0.7, max(bar, 1), stroke=0, fill=1)
        c.setFillColor(HexColor("#" + colours["muted"]))
        c.drawCentredString(x0 + i * gap + gap / 2, y0 + 0.15 * pt, r["labels"][i])
